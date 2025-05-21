from openpyxl import load_workbook
import time
import sys
from tabulate import tabulate

def average(a: list|tuple) -> float:
    try:
        return sum(a)/len(a)
    except ZeroDivisionError:
        return -1
    
def unitary_average(a, b, req):
    lower = a[0]
    upper = b[0]
    out = [0 for x in a]
    for i in range(len(a)):
        out[i] = ((req-lower)*b[i] + (upper-req)*a[i])/(upper-lower)
    return out

def battery_percent(volt):
    a = int(123*(1-(1/((1+((volt/3.7)**80))**0.165))))
    return 100 if a>100 else 0 if a<0 else a

def setup():
    global all_tstamps, fc_data, thr_data, config, start_line, tm_type, file_timestamps, file_throttlestamps, file_tiny_summary, fc_log_file, thr_log_file, is_available, to_continue
    file_timestamps = "I:/My Drive/AssailingFalcons/Test Results/TM_Timestamps.xlsx"
    file_throttlestamps = "I:/My Drive/AssailingFalcons/Test Results/TM_Throttles.xlsx"
    file_speeds = "I:/My Drive/AssailingFalcons/Test Results/TM_Speeds.xlsx"
    file_tiny_summary = "I:/My Drive/AssailingFalcons/Test Results/New_TM.xlsx"
    all_tstamps = {}
    with open("config.txt", 'r') as file:
        conf = [x.strip() for x in file.readlines()]
        start_line = int(conf[2]) - 1
        fc_log_file = conf[7]
        thr_log_file = conf[8]
    with open("FC_logs/list.txt", 'r') as file:
        try:
            fc_log_file = file.readlines()[int(fc_log_file)-1].strip()
        except Exception:
            print("No FC Logs remaining.")
            is_available = False
    with open("Thrust_logs/list.txt", 'r') as file:
        try:
            thr_log_file = file.readlines()[int(thr_log_file)-1].strip()
        except Exception:
            print("No Thrust Logs remaining.")
            is_available = False
    if not is_available:
        return
    print(fc_log_file, thr_log_file, sep='\n')
    combine_files()
    # data[:] = data[start_line : data.index('', start_line+1)]
    # process_length = len(data) + 1
    # data[:] = [x for x in data if x]
    config = fc_data[0].strip().split(',')
    test_date = time.strftime("%d-%b-%Y", time.strptime(config[0].strip(), "%Y-%m-%d %H:%M:%S"))
    prop = config[1].strip()
    motor_kv = config[2].strip()
    esc = config[3].strip()
    batt = config[4].strip()
    config[:] = [test_date, prop, motor_kv, esc, batt]
    fc_data[:] = fc_data[1:]
    print(f"{motor_kv}\n{esc}\n{batt}\n{prop}")
    try:
        tm_type = input("Static or Dynamic? ").strip()
        if tm_type in 'nN':
            print("Cancelled.")
            sys.exit(0)
    except KeyboardInterrupt:
        print("Cancelled.")
        sys.exit(0)

def combine_files():
    global fc_data, thr_data
    with open(fc_log_file, 'r') as file:
        fc_data = [x.strip() for x in file.readlines()]
    with open(thr_log_file, 'r') as file:
        thr_data = [x.strip() for x in file.readlines()]
    fc_start = time.strptime(fc_data[1].split(',')[0], "%Y-%m-%d %H:%M:%S")
    thr_start = time.strptime(thr_data[0].split(',')[0], "%Y-%m-%d %H:%M:%S")
    # print(len(fc_data), len(thr_data))
    if fc_start > thr_start:
        thr_start_t = 0
        for i, line in enumerate(thr_data):
            if line.split(',')[0] == time.strftime("%Y-%m-%d %H:%M:%S", fc_start):
                thr_start_t = i
                break
        thr_data[:] = thr_data[thr_start_t:]
    elif thr_start > fc_start:
        fc_start_t = 1
        for i, line in enumerate(fc_data[1:]):
            if line.split(',')[0] == time.strftime("%Y-%m-%d %H:%M:%S", thr_start):
                fc_start_t = i
                break
        fc_data[:] = fc_data[0:1] + fc_data[fc_start_t-1:]
    # print(len(fc_data), len(thr_data))
    fc_end_i = fc_data.index('', 1)
    thr_end_i = thr_data.index('', 1)
    # print(fc_data[fc_end_i-1].split(','))
    # print(thr_data[thr_end_i-1].split(','))
    fc_end = time.strptime(fc_data[fc_end_i-1].split(',')[0], "%Y-%m-%d %H:%M:%S")
    thr_end = time.strptime(thr_data[thr_end_i-1].split(',')[0], "%Y-%m-%d %H:%M:%S")
    if fc_end > thr_end:
        fc_end_t = -1
        for i, line in enumerate(fc_data[::-1]):
            if line.split(',')[0] == time.strftime("%Y-%m-%d %H:%M:%S", thr_end):
                fc_end_t = len(fc_data) - i
                break
        fc_data[:] = fc_data[:fc_end_t]
    elif thr_end > fc_end:
        thr_end_t = -1
        for i, line in enumerate(thr_data[::-1]):
            if line.split(',')[0] == time.strftime("%Y-%m-%d %H:%M:%S", fc_end):
                thr_end_t = len(thr_data) - i
                break
        thr_data[:] = thr_data[:thr_end_t]
    fc_data[:] = [x for x in fc_data if x]
    thr_data[:] = [x for x in thr_data if x]
    # print(len(fc_data), len(thr_data))
    # print(fc_data[0])
    # print(fc_data[1])
    # print(thr_data[0])
    # print(fc_data[-1])
    # print(thr_data[-1])

def summarise_data():
    global final_tstamps, final_throttles, start_volt, end_volt, min_volt, has_airspeed
    last_throttle = -1
    start_volt = end_volt = min_volt = 0
    start_charge = 0
    for line in fc_data:
        if not line: continue
        values = line.split(',')
        if values[1].split()[0] in ["APC", "EMAX", "GemFan", "Generic", "HQ", "Orange", "GF"]: continue
        if values[0] not in all_tstamps.keys():
            all_tstamps[values[0]] = {"Throttle":[], "Thrust":[], "Current":[], "Current2":[], "Volt":[], "Volt2":[], "RPM1":[], "RPM2":[], "Capacity":[], "Airspeed":[]}
        all_tstamps[values[0]]["Throttle"].append(int(values[1]))
        # all_tstamps[values[0]]["Thrust"].append(float(values[2]))
        all_tstamps[values[0]]["Current"].append(float(values[2]))
        all_tstamps[values[0]]["Current2"].append(float(values[3]))
        all_tstamps[values[0]]["Volt"].append(float(values[4]))
        all_tstamps[values[0]]["Volt2"].append(float(values[5]))
        all_tstamps[values[0]]["RPM1"].append(int(values[6]))
        all_tstamps[values[0]]["RPM2"].append(int(values[7]))
        all_tstamps[values[0]]["Capacity"].append(int(values[8]))
        if len(values)==10:
            all_tstamps[values[0]]["Airspeed"].append(float(values[9]))
            has_airspeed = True
        else:
            has_airspeed = False

    for line in thr_data:
        if not line: continue
        values = line.split(',')
        if values[0] not in all_tstamps.keys():
            print("Something is seriously wrong with your data.")
        else:
            all_tstamps[values[0]]["Thrust"].append(float(values[1]))

    thrust_tstamps = {}
    for line in thr_data:
        values = [x.strip() for x in line.split(',')]
        if values[0] not in thrust_tstamps.keys():
            thrust_tstamps[values[0]] = []
        else:
            thrust_tstamps[values[0]].append(int(values[1]))

    # update averaged values for timestamps
    final_tstamps = {}
    for timestamp in all_tstamps:
        if len(final_tstamps) == 0:
            start_charge = min(all_tstamps[timestamp]["Capacity"])
        if timestamp not in final_tstamps.keys():
            final_tstamps[timestamp] = {"Throttle":[], "Thrust":[], "Current":[], "Current2":[], "Volt":[], "Volt2":[], "RPM1":[], "RPM2": [], "Capacity":[]}
        final_tstamps[timestamp]["Throttle"] = int(average(all_tstamps[timestamp]["Throttle"]))
        final_tstamps[timestamp]["Thrust"] = average(all_tstamps[timestamp]["Thrust"])
        final_tstamps[timestamp]["Current"] = average(all_tstamps[timestamp]["Current"])
        final_tstamps[timestamp]["Current2"] = average(all_tstamps[timestamp]["Current2"])
        final_tstamps[timestamp]["Volt"] = average(all_tstamps[timestamp]["Volt"])
        final_tstamps[timestamp]["Volt2"] = average(all_tstamps[timestamp]["Volt2"])
        final_tstamps[timestamp]["RPM1"] = int(average(all_tstamps[timestamp]["RPM1"]))
        final_tstamps[timestamp]["RPM2"] = int(average(all_tstamps[timestamp]["RPM2"]))
        final_tstamps[timestamp]["Capacity"] = int(average(all_tstamps[timestamp]["Capacity"])) - start_charge
        final_tstamps[timestamp]["Airspeed"] = average(all_tstamps[timestamp]["Airspeed"]) if has_airspeed else 0
    
    if tm_type[0] in 'Ss':
        # update averaged values for throttle
        final_thr = []
        for timestamp in thrust_tstamps:
            m = len(thrust_tstamps[timestamp])
            n = len(all_tstamps[timestamp]["Throttle"])
            if not n: continue
            li = []
            for i in range(n):
                li.append(average(thrust_tstamps[timestamp][:]))
                # final_thr.append([all_tstamps[timestamp]["Throttle"][i], average(thrust_tstamps[timestamp][(i+m)//n:((i+m)//n)+1])])
                final_thr.append([
                    all_tstamps[timestamp]["Throttle"][i],
                    int(average(thrust_tstamps[timestamp][(i*m)//n:((i+1)*m)//n])),
                    all_tstamps[timestamp]["Current"][i],
                    all_tstamps[timestamp]["Current2"][i],
                    all_tstamps[timestamp]["Volt"][i],
                    all_tstamps[timestamp]["Volt2"][i],
                    all_tstamps[timestamp]["RPM1"][i],
                    all_tstamps[timestamp]["Capacity"][i] - start_charge,
                    all_tstamps[timestamp]["Airspeed"][i] if has_airspeed else 0,
                ])
        final_throttles = []
        ls = {}
        for values in final_thr:
            # print(last_throttle, values[0])
            if last_throttle == values[0]:
                ls["Thrust"].append(values[1])
                ls["Current"].append(values[2])
                ls["Current2"].append(values[3])
                ls["Volt"].append(values[4])
                ls["Volt2"].append(values[5])
                ls["RPM1"].append(values[6])
                ls["Capacity"].append(values[7])
                ls["Airspeed"].append(values[8])
            else:
                if last_throttle==-1: pass
                else:
                    final_throttles.append([
                        last_throttle,
                        int(average(ls["Thrust"])),
                        average(ls["Current"]),
                        average(ls["Current2"]),
                        average(ls["Volt"]),
                        average(ls["Volt2"]),
                        int(average(ls["RPM1"])),
                        int(average(ls["Capacity"])),
                        average(ls["Airspeed"]) if has_airspeed else 0,
                    ])
                last_throttle = values[0]
                ls["Thrust"] = [values[1]]
                ls["Current"] = [values[2]]
                ls["Current2"] = [values[3]]
                ls["Volt"] = [values[4]]
                ls["Volt2"] = [values[5]]
                ls["RPM1"] = [values[6]]
                ls["Capacity"] = [values[7]]
                ls["Airspeed"] = [values[8]]
        final_throttles.append([0, 0, 0, 0, round(average(ls["Volt"]), 2), round(average(ls["Volt2"]), 2), 0, int(average(ls["Capacity"]))])
        start_volt = final_thr[0][4]
        end_volt = final_thr[-1][4]
        min_volt = min([x[4] for x in final_thr])
    elif tm_type[0] in 'Dd':
        pass

def summarise_further():
    global preview, make_tiny_summary
    preview = []
    preview.append(final_throttles[0])
    search = [x[0] for x in final_throttles]
    if max(search)<90:
        make_tiny_summary = False
        return
    make_tiny_summary = True
    if tm_type[0] in 'Ss':
        # 20-80
        for i in range(1, 5):
            if i*20 in search:
                preview.append(final_throttles[search.index(i*20)])
            else:
                j = 0
                low_ind = -1
                while low_ind==-1:
                    j += 1
                    try:
                        low_ind = search.index(i*20-j)
                    except ValueError:
                        low_ind = -1
                preview.append(unitary_average(final_throttles[low_ind], final_throttles[low_ind+1], i*20))
        if 100 in search:
            preview.append(final_throttles[search.index(100)])
        elif 99 in search:
            preview.append(final_throttles[search.index(99)])
        elif 98 in search:
            preview.append(final_throttles[search.index(98)])
        else:
            preview.append(final_throttles[search.index(max(search))])
        hdrs = ["Throttle", "Thrust", "Current", "Current2", "Volt", "Volt2", "RPM1", "Capacity",] + (["Airspeed"] if has_airspeed else [])
        print(tabulate(preview, headers=hdrs, tablefmt="simple_grid"))

    elif tm_type[0] in 'Dd':
        # use final_speeds
        # (only int speeds)
        
        # 5-max([x[-1] for x in final_speeds])
        pass

def write_to_files():
    global is_available, to_continue
    name = tm_type[0] + config[0]
    wb1 = load_workbook(file_timestamps)
    if name in wb1:
        ws1 = wb1[name]
    else:
        ws1 = wb1.create_sheet(name)
    wb1.active = ws1
    start_row = ws1.max_row + 3
    if start_row==4:
        start_row = 2
    ws1.cell(row=start_row+0, column=2).value = f"{config[2]}"
    ws1.merge_cells(start_column=2, start_row=start_row, end_column=12, end_row=start_row)
    ws1.cell(row=start_row+1, column=2).value = f"({tm_type}) {config[0]}"
    ws1.merge_cells(start_column=2, start_row=start_row+1, end_column=12, end_row=start_row+1)
    ws1.cell(row=start_row+2, column=2).value = f"{config[1]}"
    ws1.merge_cells(start_column=2, start_row=start_row+2, end_column=12, end_row=start_row+2)
    ws1.cell(row=start_row+3, column=2).value = f"{config[3]}"
    ws1.merge_cells(start_column=2, start_row=start_row+3, end_column=12, end_row=start_row+3)
    ws1.cell(row=start_row+4, column=2).value = f"{config[4]}"
    ws1.merge_cells(start_column=2, start_row=start_row+4, end_column=12, end_row=start_row+4)
    batt_data = f"{battery_percent(start_volt/4)}% {start_volt:.2f}V  ->  {battery_percent(end_volt/4)}% {end_volt:.2f}V  (100% = {battery_percent(min_volt/4)}% {min_volt:.2f}V)"
    print(batt_data)
    ws1.cell(row=start_row+5, column=2).value = batt_data
    ws1.merge_cells(start_column=2, start_row=start_row+5, end_column=12, end_row=start_row+5)
    ws1.merge_cells(start_column=2, start_row=start_row+6, end_column=12, end_row=start_row+6)
    ws1.cell(row=start_row+7, column=2).value = "Timestamp"
    ws1.cell(row=start_row+7, column=3).value = "Throttle (%)"
    ws1.cell(row=start_row+7, column=4).value = "Thrust (gf)"
    ws1.cell(row=start_row+7, column=5).value = "Current (A)"
    ws1.cell(row=start_row+7, column=6).value = "Power (W)"
    ws1.cell(row=start_row+7, column=7).value = "Volt (V)"
    ws1.cell(row=start_row+7, column=8).value = "Current2 (A)"
    ws1.cell(row=start_row+7, column=9).value = "Volt2 (V)"
    ws1.cell(row=start_row+7, column=10).value = "RPM1"
    ws1.cell(row=start_row+7, column=11).value = "Capacity (mAh)"
    ws1.cell(row=start_row+7, column=12).value = "Thrust/Current"
    if has_airspeed:
        ws1.cell(row=start_row+7, column=13).value = "Airspeed (m/s)"
    start_row = start_row + 8
    wb1.save(file_timestamps)
    
    for i, tstamp in enumerate(final_tstamps.keys()):
        ws1.cell(row=start_row+i, column=2).value = tstamp
        ws1.cell(row=start_row+i, column=3).value = int(final_tstamps[tstamp]["Throttle"])
        ws1.cell(row=start_row+i, column=4).value = int(final_tstamps[tstamp]["Thrust"])
        ws1.cell(row=start_row+i, column=5).value = round(final_tstamps[tstamp]["Current"], 2)
        ws1.cell(row=start_row+i, column=6).value = round(final_tstamps[tstamp]["Current"]*final_tstamps[tstamp]["Volt"], 2)
        ws1.cell(row=start_row+i, column=7).value = round(final_tstamps[tstamp]["Volt"], 2)
        ws1.cell(row=start_row+i, column=8).value = round(final_tstamps[tstamp]["Current2"], 2)
        ws1.cell(row=start_row+i, column=9).value = round(final_tstamps[tstamp]["Volt2"], 2)
        ws1.cell(row=start_row+i, column=10).value = int(final_tstamps[tstamp]["RPM1"])
        ws1.cell(row=start_row+i, column=11).value = int(final_tstamps[tstamp]["Capacity"])
        ws1.cell(row=start_row+i, column=12).value = round(final_tstamps[tstamp]["Thrust"]/final_tstamps[tstamp]["Current"], 6) if final_tstamps[tstamp]["Current"] else 0
        if has_airspeed:
            ws1.cell(row=start_row+i, column=13).value = round(final_tstamps[tstamp]["Airspeed"], 2)
    wb1.save(file_timestamps)
    wb1.close()
    print("Timestamped data saved to file.")


    # if stm
    wb2 = load_workbook(file_throttlestamps)
    if name in wb2:
        ws2 = wb2[name]
    else:
        ws2 = wb2.create_sheet(name)
    wb2.active = ws2
    start_row = ws2.max_row + 3
    if start_row==4:
        start_row = 2
    ws2.cell(row=start_row+0, column=2).value = f"{config[2]}"
    ws2.merge_cells(start_column=2, start_row=start_row, end_column=11, end_row=start_row)
    ws2.cell(row=start_row+1, column=2).value = f"({tm_type}) {config[0]}"
    ws2.merge_cells(start_column=2, start_row=start_row+1, end_column=11, end_row=start_row+1)
    ws2.cell(row=start_row+2, column=2).value = f"{config[1]}"
    ws2.merge_cells(start_column=2, start_row=start_row+2, end_column=11, end_row=start_row+2)
    ws2.cell(row=start_row+3, column=2).value = f"{config[3]}"
    ws2.merge_cells(start_column=2, start_row=start_row+3, end_column=11, end_row=start_row+3)
    ws2.cell(row=start_row+4, column=2).value = f"{config[4]}"
    ws2.merge_cells(start_column=2, start_row=start_row+4, end_column=11, end_row=start_row+4)
    ws2.cell(row=start_row+5, column=2).value = batt_data
    ws2.merge_cells(start_column=2, start_row=start_row+5, end_column=11, end_row=start_row+5)
    ws2.merge_cells(start_column=2, start_row=start_row+6, end_column=11, end_row=start_row+6)
    ws2.cell(row=start_row+7, column=2).value = "Throttle (%)"
    ws2.cell(row=start_row+7, column=3).value = "Thrust (gf)"
    ws2.cell(row=start_row+7, column=4).value = "Current (A)"
    ws2.cell(row=start_row+7, column=5).value = "Power (W)"
    ws2.cell(row=start_row+7, column=6).value = "Volt (V)"
    ws2.cell(row=start_row+7, column=7).value = "Current2 (A)"
    ws2.cell(row=start_row+7, column=8).value = "Volt2 (V)"
    ws2.cell(row=start_row+7, column=9).value = "RPM1"
    ws2.cell(row=start_row+7, column=10).value = "Capacity (mAh)"
    ws2.cell(row=start_row+7, column=11).value = "Thrust/Current"
    if has_airspeed:
        ws2.cell(row=start_row+7, column=12).value = "Airspeed (m/s)"
    start_row = start_row + 8
    wb2.save(file_throttlestamps)

    for i, tstamp in enumerate(final_throttles):
        ws2.cell(row=start_row+i, column=2).value = int(tstamp[0])
        ws2.cell(row=start_row+i, column=3).value = int(tstamp[1])
        ws2.cell(row=start_row+i, column=4).value = round(tstamp[2], 2)
        ws2.cell(row=start_row+i, column=5).value = round(tstamp[2]*tstamp[4], 2)
        ws2.cell(row=start_row+i, column=6).value = round(tstamp[4], 2)
        ws2.cell(row=start_row+i, column=7).value = round(tstamp[3], 2)
        ws2.cell(row=start_row+i, column=8).value = round(tstamp[5], 2)
        ws2.cell(row=start_row+i, column=9).value = int(tstamp[6])
        ws2.cell(row=start_row+i, column=10).value = int(tstamp[7])
        ws2.cell(row=start_row+i, column=11).value = round(tstamp[1]/tstamp[2], 6) if tstamp[2] else 0
        if has_airspeed:
            ws2.cell(row=start_row+i, column=12).value = round(tstamp[8], 2)
    wb2.save(file_throttlestamps)
    wb2.close()
    print("Throttle-stamped data saved to file.")

    # if dtm
    # wb3 = load_workbook(file_speeds)
    # add speeds in km/h

def write_tiny_summary():
    name = tm_type[0] + config[0]
    wb = load_workbook(file_tiny_summary)
    if name in wb:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
    wb.active = ws
    start_row = ws.max_row + 3
    if start_row==4:
        start_row = 2
    ws.cell(row=start_row+0, column=2).value = f"{config[2]}"
    ws.merge_cells(start_column=2, start_row=start_row, end_column=6, end_row=start_row)
    ws.cell(row=start_row+1, column=2).value = f"({tm_type}) {config[0]}"
    ws.merge_cells(start_column=2, start_row=start_row+1, end_column=6, end_row=start_row+1)
    ws.cell(row=start_row+2, column=2).value = f"{config[1]}"
    ws.merge_cells(start_column=2, start_row=start_row+2, end_column=6, end_row=start_row+2)
    ws.cell(row=start_row+3, column=2).value = f"{config[3]}"
    ws.merge_cells(start_column=2, start_row=start_row+3, end_column=6, end_row=start_row+3)
    ws.cell(row=start_row+4, column=2).value = f"{config[4]}"
    ws.merge_cells(start_column=2, start_row=start_row+4, end_column=6, end_row=start_row+4)
    ws.cell(row=start_row+5, column=2).value = f"{battery_percent(start_volt/4)}% {start_volt:.2f}V  ->  {battery_percent(end_volt/4)}% {end_volt:.2f}V  (100% = {battery_percent(min_volt/4)}% {min_volt:.2f}V)"
    ws.merge_cells(start_column=2, start_row=start_row+5, end_column=6, end_row=start_row+5)
    ws.merge_cells(start_column=2, start_row=start_row+6, end_column=6, end_row=start_row+6)
    ws.cell(row=start_row+7, column=2).value = "Throttle (%)"
    ws.cell(row=start_row+7, column=3).value = "Thrust (gf)"
    ws.cell(row=start_row+7, column=4).value = "Current (A)"
    ws.cell(row=start_row+7, column=5).value = "Power (W)"
    ws.cell(row=start_row+7, column=6).value = "Volt (V)"
    start_row = start_row + 8
    wb.save(file_tiny_summary)
    # throttle, thrust, current, current2, volt, volt2, rpm1, capacity
    for i, tval in enumerate(preview):
        ws.cell(row=start_row+i, column=2).value = int(tval[0])
        ws.cell(row=start_row+i, column=3).value = int(tval[1])
        ws.cell(row=start_row+i, column=4).value = round(tval[2], 2)
        ws.cell(row=start_row+i, column=5).value = round(tval[2]*tval[4], 2)
        ws.cell(row=start_row+i, column=6).value = round(tval[4], 2)
    wb.save(file_tiny_summary)
    wb.close()
    print("Tiny summary saved to file.")

def update_config():
    global to_continue, is_available
    conf = fc = th = ""
    with open("config.txt", 'r') as file:
        conf = [x.strip() for x in file.readlines()]
        fc = int(conf[7]) + 1
        th = int(conf[8]) + 1
    with open("config.txt", 'w') as file:
        print('\n'.join(conf[:-2]), file=file)
        print(fc, file=file)
        print(th, file=file)
        print("Configuration file updated.")
    with open("FC_logs/list.txt", 'r') as file:
        fclen = len([x.strip() for x in file.readlines() if x.strip()])
    with open("Thrust_logs/list.txt", 'r') as file:
        thlen = len([x.strip() for x in file.readlines() if x.strip()])
    if (fclen<fc) and (thlen<th):
        to_continue = False
        is_available = False
    else:
        print(f"{fclen-fc+1} log{'' if (fclen-fc+1)==1 else 's'} remaining.")

if __name__ == "__main__":
    is_available = True
    to_continue = True
    while is_available and to_continue:
        print("Setting up...")
        setup()
        if not is_available:
            continue
        print()
        print("Setup complete.")
        summarise_data()
        write_to_files()
        summarise_further()
        if make_tiny_summary:
            write_tiny_summary()
        update_config()
        if is_available:
            s = input("Continue? y/n: ")
            if s in 'Yy':
                to_continue = True
                print()
            else: to_continue = False
        else: to_continue = False
    if not is_available:
        print("All logs processed.")
